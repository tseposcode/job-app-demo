"""Make three made-up CVs in real-world styles, to test the CV reader on CVs it
was not written for: no "LOOKING FOR" section, different headings and layouts,
and one Word file.

Run: python3 samples/make_real_world_cvs.py   (needs: pip install reportlab python-docx)

  Nomsa_Dlamini_CV.pdf     one column, "CURRICULUM VITAE" title, Title Case headings,
                           labelled personal details (with a dummy ID number line)
  Kabelo_Mahlangu_CV.pdf   two columns: a shaded sidebar like a Canva template
  Zanele_Khumalo_CV.docx   Word file with a personal details table and bullet lists

These people are fictional. Phone numbers use 000, emails use example.com,
the ID number is all zeros, and companies and schools are marked made-up.
"""

from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / 'cvs'
INK = HexColor('#142019')
MUTED = HexColor('#5B6A60')
ACCENT = HexColor('#0F6B4F')
LIGHT = HexColor('#E8EFEA')
NOTE = 'Sample CV for a Job Auto demo. This person is made up.'


def wrap(text, max_chars):
    words, line, out = text.split(), '', []
    for w in words:
        if len(line) + len(w) + 1 > max_chars:
            out.append(line)
            line = w
        else:
            line = f'{line} {w}'.strip()
    out.append(line)
    return out


def nomsa_pdf():
    path = OUT / 'Nomsa_Dlamini_CV.pdf'
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle('Nomsa Dlamini CV (sample)')
    width, height = A4
    x, y = 56, height - 64

    def text(t, size=10, font='Helvetica', dx=0, gap=14, colour=INK):
        nonlocal y
        c.setFillColor(colour)
        c.setFont(font, size)
        c.drawString(x + dx, y, t)
        y -= gap

    def heading(t):
        nonlocal y
        y -= 6
        text(t, 12, 'Helvetica-Bold', gap=6, colour=ACCENT)
        c.setStrokeColor(LIGHT)
        c.line(x, y, width - x, y)
        y -= 14

    text('CURRICULUM VITAE', 9, 'Helvetica-Bold', gap=22, colour=MUTED)
    text('Nomsa Dlamini', 22, 'Helvetica-Bold', gap=16)
    text(NOTE, 8, 'Helvetica-Oblique', gap=8, colour=MUTED)
    heading('Personal Details')
    for label, value in [
        ('Surname', 'Dlamini'),
        ('First names', 'Nomsa Precious'),
        ('ID number', '0000000000000'),
        ('Cell no', '082 000 4567'),
        ('Email', 'nomsa.dlamini@example.com'),
        ('Address', '15 Example Street, Hatfield, Pretoria'),
        ("Driver's licence", 'Code 8'),
    ]:
        c.setFont('Helvetica', 10)
        c.setFillColor(MUTED)
        c.drawString(x, y, f'{label}:')
        c.setFillColor(INK)
        c.drawString(x + 110, y, value)
        y -= 14
    heading('Career Objective')
    for line in wrap(
        'Human Resource Management graduate looking for a graduate programme or internship where I can '
        'grow in recruitment and employee relations. I am willing to relocate for the right opportunity.',
        100,
    ):
        text(line)
    heading('Education')
    text('Bachelor of Commerce in Human Resource Management', font='Helvetica-Bold')
    text('University of Johannesburg, 2022 - 2025')
    y -= 4
    text('National Senior Certificate (Matric)', font='Helvetica-Bold')
    text('Hatfield Example Secondary School (made-up), 2021')
    heading('Work Experience')
    text('HR Administration Intern | Example Staffing (made-up company) | Jan 2025 - Jun 2025', font='Helvetica-Bold')
    for b in [
        'Captured new employee records and filed employment contracts.',
        'Helped schedule interviews and onboarding sessions for 30 new staff.',
    ]:
        text(f'•  {b}', dx=10)
    text('Admin Assistant (vacation work) | Example Municipal Office (made-up) | Dec 2023 - Jan 2024', font='Helvetica-Bold')
    text('•  Answered queries, filed documents and updated spreadsheets.', dx=10)
    heading('Skills')
    for s in [
        'MS Excel (intermediate)',
        'Recruitment and onboarding',
        'Microsoft Office',
        'Good communication skills',
        'Attention to detail',
    ]:
        text(f'•  {s}', dx=10)
    heading('Languages')
    text('English, isiZulu, Sepedi')
    heading('References')
    text('Available on request')
    c.showPage()
    c.save()
    return path


def kabelo_pdf():
    path = OUT / 'Kabelo_Mahlangu_CV.pdf'
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle('Kabelo Mahlangu CV (sample)')
    width, height = A4
    side_w = 190
    c.setFillColor(LIGHT)
    c.rect(0, 0, side_w, height, stroke=0, fill=1)

    sx, sy = 28, height - 150

    def side(t, size=9.5, font='Helvetica', gap=14, colour=INK):
        nonlocal sy
        c.setFillColor(colour)
        c.setFont(font, size)
        c.drawString(sx, sy, t)
        sy -= gap

    side('CONTACT', 11, 'Helvetica-Bold', gap=18, colour=ACCENT)
    side('083 000 3579')
    side('kabelo.mahlangu@example.com', 8.5)
    side('Polokwane, Limpopo', gap=26)
    side('SKILLS', 11, 'Helvetica-Bold', gap=18, colour=ACCENT)
    for s in ['AutoCAD', 'Site inspection', 'Health and safety', 'Problem solving', 'MS Excel', 'Python (basic)']:
        side(f'• {s}')
    sy -= 12
    side('LANGUAGES', 11, 'Helvetica-Bold', gap=18, colour=ACCENT)
    for s in ['English', 'Sepedi', 'Afrikaans']:
        side(s)

    mx, my = side_w + 30, height - 70

    def main(t, size=10, font='Helvetica', gap=14, colour=INK):
        nonlocal my
        c.setFillColor(colour)
        c.setFont(font, size)
        c.drawString(mx, my, t)
        my -= gap

    def mhead(t):
        nonlocal my
        my -= 8
        main(t, 13, 'Helvetica-Bold', gap=18, colour=ACCENT)

    main('KABELO MAHLANGU', 24, 'Helvetica-Bold', gap=20)
    main('Civil Engineering Graduate', 11, gap=12, colour=MUTED)
    main(NOTE, 7.5, 'Helvetica-Oblique', gap=10, colour=MUTED)
    mhead('Profile')
    for line in wrap(
        'Civil engineering graduate with site experience on road and bridge projects. '
        'Careful, safety-minded and quick to learn. Open to relocating anywhere in South Africa.',
        62,
    ):
        main(line)
    mhead('Experience')
    main('Vacation Work Student', font='Helvetica-Bold')
    main('Example Roads & Bridges (made-up company)', colour=MUTED)
    main('Dec 2024 - Jan 2025', 9, colour=MUTED)
    for b in ['Checked concrete pours and kept daily site diaries.', 'Drew small site layouts in AutoCAD.']:
        main(f'• {b}')
    my -= 6
    main('Student Assistant', font='Helvetica-Bold')
    main('Tshwane University of Technology, Civil Engineering', colour=MUTED)
    main('2023 - 2024', 9, colour=MUTED)
    main('• Marked first-year drawing tutorials.')
    mhead('Education')
    main('Bachelor of Engineering Technology in Civil Engineering', font='Helvetica-Bold')
    main('Tshwane University of Technology | 2022 - 2025', colour=MUTED)
    my -= 6
    main('National Senior Certificate', font='Helvetica-Bold')
    main('Seshego Example High School (made-up) | 2021', colour=MUTED)
    c.showPage()
    c.save()
    return path


def zanele_docx():
    from docx import Document
    from docx.shared import Pt

    path = OUT / 'Zanele_Khumalo_CV.docx'
    doc = Document()
    doc.core_properties.title = 'Zanele Khumalo CV (sample)'
    doc.add_heading('ZANELE KHUMALO', level=0)
    doc.add_paragraph('Marketing Graduate')
    note = doc.add_paragraph(NOTE)
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(8)

    def head(t):
        r = doc.add_paragraph().add_run(t)
        r.bold = True
        r.font.size = Pt(12)

    head('PERSONAL DETAILS')
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in [
        ('Full name', 'Zanele Khumalo'),
        ('Cell number', '076 000 8642'),
        ('Email', 'zanele.khumalo@example.com'),
        ('Location', 'Umlazi, KZN'),
        ('Health', 'Good'),
        ("Driver's licence", 'No'),
    ]:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
    head('CAREER OBJECTIVE')
    doc.add_paragraph(
        'Creative marketing graduate who loves building brands online. Looking for a learnership or '
        'internship in digital marketing. I am also open to remote work.'
    )
    head('QUALIFICATIONS')
    doc.add_paragraph(
        'Advanced Diploma in Marketing Management - Mangosuthu University of Technology (in progress, completing 2026)'
    )
    doc.add_paragraph('Diploma in Marketing - Mangosuthu University of Technology, 2024')
    doc.add_paragraph('National Senior Certificate - Umlazi Example High School (made-up), 2020')
    head('SKILLS')
    for s in ['Social media management', 'Canva', 'Copywriting', 'Customer service', 'MS Office']:
        doc.add_paragraph(s, style='List Bullet')
    head('WORK EXPERIENCE')
    doc.add_paragraph('Social Media Assistant (part-time), Example Spaza Brands (made-up company), 2024 - present')
    for b in [
        'Planned and posted content on Instagram and Facebook three times a week.',
        'Grew the page from 400 to 2 100 followers in eight months.',
    ]:
        doc.add_paragraph(b, style='List Bullet')
    head('REFERENCES')
    doc.add_paragraph('Available on request')
    doc.save(str(path))
    return path


if __name__ == '__main__':
    OUT.mkdir(exist_ok=True)
    for make in (nomsa_pdf, kabelo_pdf, zanele_docx):
        print('made', make())
